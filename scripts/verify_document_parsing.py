"""Opt-in live document vision check using synthetic PNG/PDF/DOCX only."""
import asyncio
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from backend.core.config import Settings
from backend.core.db import Database
from backend.infrastructure.providers import Providers
from backend.services.documents import Documents


async def main():
    folder = Path('.cache/document-smoke')
    folder.mkdir(parents=True, exist_ok=True)
    picture = Image.new('RGB', (1100, 500), 'white')
    draw = ImageDraw.Draw(picture)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    font = ImageFont.truetype(str(font_path), 30) if font_path.exists() else ImageFont.load_default(size=30)
    draw.text((30, 20), 'SYNTHETIC TEST DATA', fill='black', font=font)
    draw.text((30, 140), 'Team A: 12 completed', fill='black', font=font)
    draw.text((570, 140), 'Team B: 7 NOT completed', fill='black', font=font)
    picture.save(folder / 'scan.png')
    picture.save(folder / 'scan.pdf', resolution=120)
    doc = Document()
    doc.add_heading('Synthetic operations', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Team', 'Status'
    row = table.add_row().cells
    row[0].text, row[1].text = 'A', '12 completed'
    row = table.add_row().cells
    row[0].text, row[1].text = 'B', '7 NOT completed'
    doc.add_paragraph('The image below is a synthetic source.')
    doc.add_picture(str(folder / 'scan.png'))
    doc.save(folder / 'mixed.docx')
    settings = Settings(demo_mode=False)
    db = Database(folder / 'usage.sqlite3')
    await db.init()
    provider = Providers(settings, db)
    responses = []
    original = provider.post
    async def recorded(*args, **kwargs):
        response = await original(*args, **kwargs)
        responses.append({'model': response.get('model'), 'usage': response.get('usage'),
                          'finish_reason': (response.get('choices') or [{}])[0].get('finish_reason')})
        return response
    provider.post = recorded
    documents = Documents(settings, db, provider, None)
    report = {'synthetic_only': True, 'requested_model': settings.vision_model_id, 'files': []}
    try:
        for name in ('scan.pdf', 'mixed.docx'):
            chunks = await documents.split_for_index(name, (folder / name).read_bytes())
            visual = '\n'.join(c['text'] for c in chunks if '图片解析' in c['text'])
            checks = {'number_12': '12' in visual, 'number_7': '7' in visual,
                      'negation': any(v in visual for v in ('NOT', '未完成', '没有完成')),
                      'locator': all(c['locator'] for c in chunks)}
            report['files'].append({'name': name, 'chunks': len(chunks), 'checks': checks})
        report['responses'] = responses
        (folder / 'result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), 'utf-8')
        print(json.dumps(report, ensure_ascii=True))
        assert all(all(f['checks'].values()) for f in report['files'])
    finally:
        await provider.close()


if __name__ == '__main__':
    asyncio.run(main())
