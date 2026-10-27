"""Direct CPython security/ABI probes inside the full production image.

Uses only the official signed-source archive and in-memory TLS; no supplier HTTP.
The deliberately distrusted certificates are public upstream regression fixtures,
not application credentials, and are never copied into the runtime image.
"""
import argparse
import ast
import bz2
import gzip
import hashlib
import importlib.metadata
import io
import json
import lzma
import os
from pathlib import Path
import platform
import sqlite3
import ssl
import sys
import tarfile
import tempfile
import xml.etree.ElementTree as ET
from xml.parsers import expat
import zipfile

if __name__ == '__main__':
    sys.path.insert(0, '/app')

SOURCE_SHA256='9120e0a653868f6e3698d36cdd99e5a58903ec9469e448f21d98f7b16d8fa290'


def handshake(client_context,server_context,hostname='localhost'):
    incoming,outgoing,server_in,server_out=(ssl.MemoryBIO() for _ in range(4))
    client=client_context.wrap_bio(incoming,outgoing,server_hostname=hostname)
    server=server_context.wrap_bio(server_in,server_out,server_side=True)
    client_done=server_done=False
    for _ in range(100):
        try:
            client.do_handshake()
            client_done=True
        except (ssl.SSLWantReadError,ssl.SSLWantWriteError):
            pass
        payload=outgoing.read()
        if payload:
            server_in.write(payload)
        try:
            server.do_handshake()
            server_done=True
        except (ssl.SSLWantReadError,ssl.SSLWantWriteError):
            pass
        payload=server_out.read()
        if payload:
            incoming.write(payload)
        if client_done and server_done:
            return client.version()
    raise AssertionError('Local TLS handshake did not finish')


def verify(archive):
    assert sys.version_info[:3]==(3,11,18)
    assert hashlib.sha256(archive.read_bytes()).hexdigest()==SOURCE_SHA256
    assert os.getuid()==10001
    assert expat.EXPAT_VERSION=='expat_2.9.0'
    with tempfile.TemporaryDirectory() as temporary:
        root=Path(temporary)
        with tarfile.open(archive) as source:
            text=source.extractfile('Python-3.11.18/Lib/test/test_ssl.py').read().decode()
            constants={}
            for node in ast.parse(text).body:
                if isinstance(node,ast.Assign) and isinstance(node.value,ast.Constant):
                    for target in node.targets:
                        if isinstance(target,ast.Name):
                            constants[target.id]=node.value.value
            blob=source.extractfile('Python-3.11.18/Lib/test/certdata/distrusted.zip').read()
            with zipfile.ZipFile(io.BytesIO(blob)) as certificates:
                certificates.extractall(root,pwd=constants['DISTRUSTED_SPKI_CHOSEN_KNIGHT'])
            for name in ('keycert3.pem','pycacert.pem'):
                (root/name).write_bytes(source.extractfile('Python-3.11.18/Lib/test/certdata/'+name).read())
        client=ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
        client.load_verify_locations(root/'distrusted/distrusted_root.pem')
        server=ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        server.load_cert_chain(root/'distrusted/distrusted_keycert.pem')
        try:
            handshake(client,server,constants['DISTRUSTED_SPKI_HOSTNAME'])
            raise AssertionError('Compromised eMudhra intermediate was trusted')
        except ssl.SSLCertVerificationError as error:
            assert error.verify_code==28
        valid_client=ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
        valid_client.load_verify_locations(root/'pycacert.pem')
        valid_server=ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        valid_server.load_cert_chain(root/'keycert3.pem')
        tls=handshake(valid_client,valid_server)
        # The exact empty/destination link case changed by CVE2026-114583.
        for name in ('','.', './'):
            for kind in (tarfile.SYMTYPE,tarfile.LNKTYPE):
                member=tarfile.TarInfo(name)
                member.type=kind
                member.linkname='dummy'
                try:
                    tarfile.tar_filter(member,str(root))
                    raise AssertionError('tar_filter accepted a link replacing its extraction root')
                except tarfile.OutsideDestinationError:
                    pass
        safe=tarfile.TarInfo('inside.txt')
        assert tarfile.tar_filter(safe,str(root)).name=='inside.txt'
    payload='native closure UTF8 sample'.encode()
    for module in (gzip,bz2,lzma):
        assert module.decompress(module.compress(payload))==payload
    with sqlite3.connect(':memory:') as database:
        assert database.execute('SELECT 42').fetchone()[0]==42
    import dbm.gnu  # noqa:F401
    import dbm.ndbm  # noqa:F401
    with tempfile.TemporaryDirectory() as temporary:
        with dbm.ndbm.open(str(Path(temporary)/'abi-fixture'),'c') as database:
            database[b'key']=b'value'
        with dbm.ndbm.open(str(Path(temporary)/'abi-fixture'),'r') as database:
            assert database[b'key']==b'value'
    assert ET.fromstring('<root>safe</root>').text=='safe'
    from docx import Document
    from backend.services.document_parsing import validate_upload,split_document
    document=Document()
    document.add_heading('Runtime parser fixture',0)
    document.add_paragraph('Signed Python source closure retains document parsing.')
    output=io.BytesIO()
    document.save(output)
    validate_upload('fixture.docx',output.getvalue())
    chunks=split_document('fixture.docx',output.getvalue())
    assert any('Signed Python source closure' in row['text'] for row in chunks)
    import backend.api.app  # noqa:F401
    import asyncpg  # noqa:F401
    import psycopg  # noqa:F401
    import pypdfium2  # noqa:F401
    import numpy  # noqa:F401
    return {'passed':True,'python_version':platform.python_version(),'openssl':ssl.OPENSSL_VERSION,
        'expat':expat.EXPAT_VERSION,'uid':os.getuid(),'tls':tls,'ca_count':len(ssl.create_default_context().get_ca_certs()),
        'checks':['CVE2026-115557 compromised intermediate denied code28 using upstream fixture','valid verified TLS handshake accepted',
        'CVE2026-114583 extraction-destination symbolic/hardlink rejected','valid tar member accepted','bundled Expat2.9.0 XML parse',
        'gzip/bz2/lzma/sqlite nativeclosure','real DOCX validation+splitter','API/asyncpg/psycopg/PDFium/numpy imports'],
        'installed_distributions':dict(sorted((d.metadata['Name'],d.version) for d in importlib.metadata.distributions())),
        'source_archive_sha256':SOURCE_SHA256,'supplier_http':0,'network_http':0,
        'scope':'Direct signed-source interpreter/security/native ABI checks only; not fullmodelquality, deploymentrole oravailability evidence.'}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-archive',required=True,type=Path)
    args=parser.parse_args()
    print(json.dumps(verify(args.source_archive),ensure_ascii=False))
