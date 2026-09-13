"""Deterministic malformed-data, layout, and retrieval degradation regressions."""
import io
import json
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from backend.infrastructure.providers import ServiceError
from backend.research.graph import evidence_context, merge_evidence, ResearchGraph
from backend.services.document_parsing import extract_blocks, split_document
from backend.services.retrieval_contract import normalize_hits
from tests.test_providers import make_provider


def pdf_bytes(*, columns=False, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    data = 'BT /F1 12 Tf 30 740 Td (Quarterly research evidence. Revenue report and scope.) Tj ET\n'
    if columns:
        data = '\n'.join(f'BT /F1 12 Tf {x} {y} Td ({text}) Tj ET' for y in (700, 680, 660) for x, text in ((30, 'LEFT Alpha 12'), (330, 'RIGHT Beta 7')))
    stream = DecodedStreamObject()
    stream.set_data(data.encode())
    page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted:
        writer.encrypt('fixture-password')
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_plain_columns_scan_and_encrypted_are_distinguished():
    plain = extract_blocks('plain.pdf', pdf_bytes())
    assert plain and plain[0].text.startswith('Quarterly') and not plain[0].image
    columns = extract_blocks('columns.pdf', pdf_bytes(columns=True))
    assert len(columns) == 1 and columns[0].image and '整页' in columns[0].locator
    assert Image.open(io.BytesIO(columns[0].image)).width <= 2000
    scanned = io.BytesIO()
    Image.new('RGB', (100, 100), 'white').save(scanned, format='PDF')
    assert extract_blocks('scan.pdf', scanned.getvalue())[0].image
    with pytest.raises(ServiceError, match='加密'):
        extract_blocks('locked.pdf', pdf_bytes(encrypted=True))
    with pytest.raises(ServiceError, match='无法解析'):
        extract_blocks('broken.pdf', b'%PDF-1.4\ncorrupt')


def test_docx_preserves_body_order_table_header_and_embedded_image():
    doc = Document()
    doc.add_heading('Operations', level=1)
    doc.add_paragraph('Before table')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Team', 'Count'
    for i in range(12):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = f'Alpha-{i}', str(i)
    doc.add_paragraph('After table')
    image = io.BytesIO()
    Image.new('RGB', (32, 32), 'white').save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image)
    output = io.BytesIO()
    doc.save(output)
    blocks = extract_blocks('layout.docx', output.getvalue())
    texts = [b.text for b in blocks]
    assert texts.index('Before table') < next(i for i, t in enumerate(texts) if 'Alpha-0' in t) < texts.index('After table')
    assert all('Team' in b.text and 'Count' in b.text for b in blocks if 'Alpha-' in b.text)
    assert blocks[-1].image and 'Operations' in blocks[-1].locator
    chunks = split_document('layout.docx', output.getvalue())
    assert all(len(c['text']) <= 1100 for c in chunks)


def test_markdown_headings_tables_and_invalid_text():
    parts = split_document('report.md', '# Quarter\n\n| Team | Count |\n| --- | --- |\n| Alpha | 12 |\n| Beta | 7 |'.encode())
    assert all('Quarter' in p['locator'] for p in parts)
    assert all('Team' in p['text'] and 'Count' in p['text'] for p in parts if 'Alpha' in p['text'] or 'Beta' in p['text'])
    with pytest.raises(ServiceError, match='列数'):
        split_document('bad.md', b'| A | B |\n| --- | --- |\n| 1 |')
    with pytest.raises(ServiceError, match='UTF-8'):
        split_document('binary.txt', b'\x00\x01\x02')


async def indexed(runtime, user='alice', text='ALPHA enterprise research facts'):
    doc = await runtime.documents.add(user, 'facts.txt', text.encode())
    await runtime.documents.ingest(doc['id'])
    return doc


async def test_bad_vector_fields_keep_keyword_evidence_without_using_remote_text(runtime, monkeypatch):
    doc = await indexed(runtime)
    row = await runtime.db.one('SELECT id FROM chunks WHERE document_id=?', (doc['id'],))
    monkeypatch.setattr(runtime.vectors, 'search', AsyncMock(return_value=[None, {'score': 1}, {'id': row['id'], 'score': 'NaN'}, {'id': 'other-workspace', 'score': .9}, {'id': row['id'], 'score': '0.9', 'text': 'INJECTED'}]))
    result = await runtime.documents.search('alice', ['ALPHA'])
    assert result and result[0]['text'] == 'ALPHA enterprise research facts'
    assert result.reasons == ['invalid_records']
    assert not await runtime.db.one('SELECT result FROM document_search_cache')
    assert not await runtime.documents.search('bob', ['ALPHA'])


async def test_vector_failure_preserves_lexical_result_and_retries_recovery(runtime, monkeypatch):
    await indexed(runtime)
    search = AsyncMock(side_effect=ServiceError('fixture outage'))
    monkeypatch.setattr(runtime.vectors, 'search', search)
    first = await runtime.documents.search('alice', ['ALPHA'])
    second = await runtime.documents.search('alice', ['ALPHA'])
    assert first == second and first.reasons == ['vector_unavailable'] and search.await_count == 2
    with pytest.raises(ServiceError, match='关键词检索也未找到'):
        await runtime.documents.search('alice', ['unrelated'])
    graph = ResearchGraph(runtime.settings, runtime.db, runtime.providers, runtime.vectors, runtime.documents)
    result = await graph.local_scout({'user_id': 'alice', 'run_id': 'rag-test', 'queries': ['ALPHA']})
    assert result['local_results'] and {'source': 'local', 'outcome': 'vector_unavailable'} in result['local_outcomes']


async def test_cache_rehydrates_provenance_and_rejects_corrupt_scores(runtime, monkeypatch):
    await indexed(runtime)
    first = await runtime.documents.search('alice', ['ALPHA'])
    cache = await runtime.db.one('SELECT result FROM document_search_cache')
    rows = json.loads(cache['result'])
    rows[0].update(text='FORGED', user_id='bob')
    await runtime.db.execute('UPDATE document_search_cache SET result=?', (json.dumps(rows),))
    assert (await runtime.documents.search('alice', ['ALPHA']))[0]['text'] == first[0]['text']
    rows[0]['score'] = None
    await runtime.db.execute('UPDATE document_search_cache SET result=?', (json.dumps(rows),))
    repaired = await runtime.documents.search('alice', ['ALPHA'])
    assert repaired and repaired.reasons == ['invalid_records']


async def test_no_zero_score_evidence_or_partial_embedding_publication(runtime, monkeypatch):
    await indexed(runtime)
    monkeypatch.setattr(runtime.vectors, 'search', AsyncMock(return_value=[]))
    assert not await runtime.documents.search('alice', ['unrelated'])
    doc = await runtime.documents.add('alice', 'new.txt', b'New source text')
    monkeypatch.setattr(runtime.providers, 'embed', AsyncMock(return_value=[]))
    with pytest.raises(ServiceError, match='数量'):
        await runtime.documents.ingest(doc['id'])
    assert (await runtime.db.one('SELECT status FROM documents WHERE id=?', (doc['id'],)))['status'] == 'failed'
    assert not await runtime.db.one('SELECT id FROM chunks WHERE document_id=?', (doc['id'],))


def test_context_cap_discloses_truncation_and_balances_origins():
    items = [dict(id=f'L{i}', text=f'{i} source.\n' * 1000, kind='local', document_id='large', access='document') for i in range(48)]
    items.append(dict(id='W1', text='external fact', kind='web', url='https://example.com', access='fulltext'))
    merged = merge_evidence(items, limit=6)
    assert 'W1' in {i['id'] for i in merged}
    context = evidence_context(merged, total_limit=6000)
    assert sum(len(i['text']) for i in context) <= 6000
    assert any(i['excerpt_truncated'] for i in context)
    assert len(items[0]['text']) > len(context[0]['text'])


async def test_document_vision_repairs_format_once_and_rejects_unreadable(tmp_path):
    calls = []
    def response(request):
        calls.append(json.loads(request.content))
        content = 'not JSON' if len(calls) == 1 else json.dumps({'readable': True, 'text': 'Alpha 12. Beta 7 NOT completed.'})
        return httpx.Response(200, json={'choices': [{'message': {'content': content}, 'finish_reason': 'stop'}]})
    provider = await make_provider(tmp_path, response)
    try:
        assert 'NOT' in await provider.describe_image(b'fixture', 'image/png')
        assert len(calls) == 2 and 'not JSON' not in json.dumps(calls[1])
    finally:
        await provider.close()


@pytest.mark.parametrize('content,finish', [(None, 'stop'), ('{"readable":true,"text":"cut"}', 'length'), ('{"readable":false,"text":"blurred"}', 'stop')])
async def test_document_vision_never_publishes_invalid_or_incomplete_extraction(tmp_path, content, finish):
    provider = await make_provider(tmp_path, lambda r: httpx.Response(200, json={'choices': [{'message': {'content': content}, 'finish_reason': finish}]}))
    try:
        with pytest.raises(ServiceError):
            await provider.describe_image(b'fixture', 'image/png')
    finally:
        await provider.close()


def test_adapter_does_not_invent_ids_or_accept_booleans():
    result, rejected = normalize_hits([{'id': 'known', 'score': True}, {'id': 'known', 'score': '0.4'}], {'known': {}}, 5)
    assert rejected and result == {'known': .4}

async def test_validator_batches_many_claims_without_losing_full_sources(runtime):
    from backend.domain.models import ReportDraft, Verification
    graph = ResearchGraph(runtime.settings, runtime.db, runtime.providers, runtime.vectors, runtime.documents)
    sources = {f'L{i}': {'id': f'L{i}', 'kind': 'local', 'text': 'Enterprise evidence. ' * 500,
                         'url': '', 'access': 'document'} for i in range(24)}
    draft = ReportDraft.model_validate({'title': 'Research', 'sections': [
        {'heading': f'Section {part}', 'claims': [{'text': '根据所提供资料，Enterprise evidence.', 'source_ids': [f'L{i}']} for i in range(part * 12, (part + 1) * 12)]} for part in range(2)]})
    seen, sizes = [], []
    async def verify(role, instruction, data, schema, run_id):
        assert role == 'validator'
        sizes.append(sum(len(s['text']) for s in data['sources'].values()))
        for source in data['sources'].values():
            assert source['text'] == 'Enterprise evidence. ' * 500
        seen.extend(c['index'] for c in data['claims'])
        return Verification(checks=[{'index': c['index'], 'supported': True, 'reason': 'supported'} for c in data['claims']])
    graph.ask = verify
    approved, _ = await graph.check_draft({'run_id': 'batch-test'}, draft, sources)
    assert approved == set(range(24)) and seen == list(range(24))
    assert len(sizes) > 1 and max(sizes) <= 96000

async def test_bad_upload_is_client_error_and_search_diagnostics_are_optional(api_client, monkeypatch):
    app, client = api_client
    invalid = await client.post('/api/documents', files={'file': ('bad.png', b'not a png', 'image/png')})
    assert invalid.status_code == 422
    await indexed(app.state.runtime)
    monkeypatch.setattr(app.state.runtime.vectors, 'search', AsyncMock(side_effect=ServiceError('offline')))
    response = await client.post('/api/documents/search?diagnostics=true', json={'query': 'ALPHA'})
    assert response.status_code == 200
    assert response.json()['results'] and response.json()['reasons'] == ['vector_unavailable']
    assert isinstance((await client.post('/api/documents/search', json={'query': 'ALPHA'})).json(), list)

async def test_corrupt_cache_does_not_turn_a_successful_empty_search_into_service_failure(runtime, monkeypatch):
    await indexed(runtime)
    monkeypatch.setattr(runtime.vectors, 'search', AsyncMock(return_value=[]))
    assert not await runtime.documents.search('alice', ['unrelated'])
    await runtime.db.execute('UPDATE document_search_cache SET result=?', ('[null]',))
    result = await runtime.documents.search('alice', ['unrelated'])
    assert not result and result.reasons == ['invalid_records']

async def test_persisted_chunk_order_survives_ingest_and_api_inspection(api_client):
    app, client = api_client
    doc = Document()
    doc.add_heading('FIRST', level=1)
    doc.add_paragraph('SECOND')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Name', 'Value'
    cells = table.add_row().cells
    cells[0].text, cells[1].text = 'THIRD', '7 NOT completed'
    doc.add_paragraph('LAST')
    content = io.BytesIO()
    doc.save(content)
    row = await app.state.runtime.documents.add('alice', 'ordered.docx', content.getvalue())
    await app.state.runtime.documents.ingest(row['id'])
    chunks = (await client.get(f'/api/documents/{row["id"]}/chunks')).json()
    assert [c['ordinal'] for c in chunks] == list(range(5))
    assert chunks[0]['text'] == 'FIRST' and chunks[-1]['text'] == 'LAST'
    assert 'THIRD' in chunks[3]['text'] and 'Name' in chunks[3]['text']
