import io
import zipfile

from docx import Document

from backend.services.documents import split_document
from backend.research.evidence_policy import validate_claim
from backend.research.graph import compact_limitations


def source(text, *, access="fulltext", domain="a.example", trust=""):
    return {"text": text, "access": access, "domain": domain, "trust_label": trust}


def test_strong_claims_need_literal_text_and_independent_sources():
    claim = "2025 年市场规模为 12%"
    assert not validate_claim(claim, [source("2025 年市场规模为 12%")])[0]
    assert validate_claim(claim, [source("2025 年市场规模为 12%", domain="a.example"),
                                  source("2025 年市场规模为 12%", domain="b.example")])[0]
    assert validate_claim(claim, [source("2025 年市场规模为 12%", trust="official")])[0]
    assert not validate_claim(claim, [source("2025 年市场规模为 10%", domain="a.example"),
                                      source("2025 年市场规模为 10%", domain="b.example")])[0]
    assert not validate_claim(claim, [source("2025 年市场规模为 12%", access="summary", domain="a.example"),
                                      source("2025 年市场规模为 12%", access="summary", domain="b.example")])[0]


def test_limitations_are_grouped_and_capped():
    limitations = compact_limitations(
        ["缺少市场规模原始统计", "缺少竞争对手披露", "缺少用户访谈"],
        ["两个来源的统计口径不一致", "更新时间不同"],
        ["一个网页正文无法访问", "一个搜索请求超时"],
        removed_claims=True,
        uses_summary=True,
    )
    assert len(limitations) == 4
    assert limitations[0].startswith("仍有 3 项研究缺口")
    assert limitations[1].startswith("存在 2 项信息口径冲突")


def test_docx_headings_and_semantic_chunks_are_preserved():
    document = Document()
    document.add_heading("市场概览", level=1)
    document.add_paragraph("企业市场研究资料。" * 120)
    buffer = io.BytesIO()
    document.save(buffer)
    chunks = split_document("report.docx", buffer.getvalue())
    assert chunks and all("市场概览" in chunk["locator"] for chunk in chunks)
    assert all(len(chunk["text"]) <= 1500 for chunk in chunks)


async def test_export_and_confirmed_data_cleanup(api_client):
    app, client = api_client
    created = (await client.post("/api/research/runs", json={"topic": "研究流程", "client_request_id": "export001"})).json()
    await app.state.runtime.tasks[created["id"]]
    assert (await client.post(f"/api/research/runs/{created['id']}/memory")).status_code == 201
    exported = await client.get("/api/data/export")
    assert exported.status_code == 200
    with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
        assert {"research.json", "memories.json", "documents.json"} <= set(archive.namelist())
    assert (await client.delete("/api/data?scope=memories")).status_code == 400
    assert (await client.delete("/api/data?scope=memories", headers={"X-Confirm-Delete": "DELETE"})).status_code == 200
    assert (await client.get("/api/memories")).json() == []


async def test_metrics_are_isolated_by_user(api_client):
    _, client = api_client
    own = (await client.get("/api/metrics")).json()
    other = (await client.get("/api/metrics", headers={"X-User-ID": "bob"})).json()
    assert own["runs"] >= 0 and other["runs"] == 0
