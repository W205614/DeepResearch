"""Format-aware extraction. Every block retains its original locator."""
import io
import re
import zipfile
import warnings
import threading
from contextlib import closing
from pathlib import Path
from dataclasses import dataclass
from PIL import Image
from docx import Document as WordDocument
from docx.table import Table
from pypdf import PdfReader
from ..infrastructure.providers import ServiceError

class InvalidDocument(ServiceError):
    """Invalid user-supplied bytes, distinct from an unavailable upstream service."""


PDFIUM_LOCK = threading.Lock()

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
ALLOWED_SUFFIXES = {".txt", ".md", ".pdf", ".docx", *IMAGE_SUFFIXES}
IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".gif": "image/gif", ".webp": "image/webp",
}


def image_media_type(content: bytes) -> str | None:
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if len(content) >= 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "image/webp"
    return None


def tokens(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9_]+|[\u4e00-\u9fff]", text.lower())


def validate_upload(name: str, content: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise InvalidDocument("仅支持 TXT、Markdown、DOCX、PDF 与 JPG/PNG/GIF/WebP 图片")
    if not content:
        raise InvalidDocument("文件内容不能为空")
    if suffix in IMAGE_SUFFIXES:
        if image_media_type(content) != IMAGE_MEDIA_TYPES[suffix]:
            raise InvalidDocument("图片文件签名无效或与扩展名不匹配")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(content)) as image:
                    if image.width * image.height > 40_000_000:
                        raise InvalidDocument("图片像素超过 4000 万限制，请缩小后上传")
                    if getattr(image, "n_frames", 1) != 1:
                        raise InvalidDocument("暂不支持动画图片，请上传静态图片")
                    image.verify()
                with Image.open(io.BytesIO(content)) as image:
                    image.load()
        except ServiceError:
            raise
        except Exception:
            raise InvalidDocument("图片损坏或无法解码，请重新导出后上传") from None
    if suffix == ".pdf" and not content.startswith(b"%PDF-"):
        raise InvalidDocument("PDF 文件签名无效")
    if suffix == ".docx":
        if not content.startswith(b"PK"):
            raise InvalidDocument("DOCX 文件签名无效")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                expanded = sum(entry.file_size for entry in entries)
                compressed = max(1, sum(entry.compress_size for entry in entries))
                if len(entries) > 200 or expanded > 40 * 1024 * 1024 or expanded / compressed > 100:
                    raise InvalidDocument("DOCX 解压后的内容超过安全限制")
                if "[Content_Types].xml" not in archive.namelist():
                    raise InvalidDocument("DOCX 文件结构无效")
        except ServiceError:
            raise
        except zipfile.BadZipFile:
            raise InvalidDocument("DOCX 文件无法解压") from None
    return suffix

def semantic_chunks(locator: str, text: str) -> list[dict]:
    text = text.strip().replace("\x00", "")
    sentences = re.split(r"(?<=[。！？!?；;])\s*", text)
    output, current, start = [], "", 1

    def flush() -> str:
        nonlocal start
        output.append({"text": current, "locator": f"{locator} · 字符 {start}–{start + len(current) - 1}"})
        overlap = current[-180:]
        start += len(current) - len(overlap)
        return overlap

    for sentence in sentences:
        if not sentence:
            continue
        while sentence:
            available = 1100 - len(current)
            if available <= 0:
                current = flush()
                continue
            if len(sentence) <= available:
                current += sentence
                break
            current += sentence[:available]
            sentence = sentence[available:]
            current = flush()
    if current:
        output.append({"text": current, "locator": f"{locator} · 字符 {start}–{start + len(current) - 1}"})
    return output



@dataclass
class Block:
    locator: str
    text: str = ''
    image: bytes = b''
    media_type: str = ''


def table_blocks(locator, rows):
    if not rows:
        return []
    header = '表头：' + ' | '.join(rows[0])
    if len(header) > 500:
        raise ServiceError(f'{locator} 表头过长，请拆分表格或转为 PDF')
    output = [Block(locator + ' · 表头', header)]
    for index, row in enumerate(rows[1:], 2):
        # Repeat the complete header in every row/continuation; never invent empty values.
        text = ' | '.join(row)
        size = 1000 - len(header)
        for offset in range(0, max(1, len(text)), size):
            output.append(Block(f'{locator} · 行 {index} · 行内字符 {offset + 1}',
                                header + '\n' + text[offset:offset + size]))
    return output


def markdown_blocks(text):
    output, heading, buffer, fenced = [], '正文', [], False
    lines = text.splitlines()
    index = 0
    def flush():
        if buffer:
            output.append(Block(heading, '\n'.join(buffer)))
            buffer.clear()
    while index < len(lines):
        line = lines[index]
        if line.lstrip().startswith(('```', '~~~')):
            fenced = not fenced
        match = re.match(r'^#{1,6}\s+(.+)', line) if not fenced else None
        if match:
            flush()
            heading = match[1][:180]
            buffer.append(line)
        elif (not fenced and '|' in line and index + 1 < len(lines)
              and re.fullmatch(r'\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*', lines[index + 1])):
            flush()
            rows = [[part.strip() for part in line.strip().strip('|').split('|')]]
            index += 2
            while index < len(lines) and '|' in lines[index] and lines[index].strip():
                row = [part.strip() for part in lines[index].strip().strip('|').split('|')]
                if len(row) != len(rows[0]):
                    raise ServiceError('Markdown 表格列数不一致，请修正分隔符或转为 PDF')
                rows.append(row)
                index += 1
            output.extend(table_blocks(heading + ' · 表格', rows))
            continue
        else:
            buffer.append(line)
        index += 1
    flush()
    return output


def docx_blocks(content):
    document = WordDocument(io.BytesIO(content))
    # These constructs are not reliably represented by python-docx's body iterator.
    unsupported = document.element.xpath('.//w:txbxContent | .//w:object | .//w:footnoteReference | .//w:endnoteReference | .//w:ins | .//w:del | .//w:altChunk | .//*[local-name()="chart"] | .//*[local-name()="relIds"]')
    if unsupported:
        raise ServiceError('DOCX 含文本框、图表对象、嵌入对象、脚注或修订内容，请接受修订并导出 PDF 后上传')
    output = []
    def walk(container, prefix='正文'):
        heading = prefix
        for index, item in enumerate(container.iter_inner_content(), 1):
            locator = f'{heading} · 段落 {index}'
            if isinstance(item, Table):
                locator = f'{heading} · 表格 {index}'
                rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                output.extend(table_blocks(locator, rows))
                seen = set()
                for r, row in enumerate(item.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        # Nested content and images need their own source location.
                        if cell.tables or cell._tc.xpath('.//a:blip'):
                            walk(cell, f'{locator} · 单元格 {r},{c}')
                continue
            if item.style and item.style.name.lower().startswith('heading') and item.text.strip():
                heading = item.text.strip()[:180]
                locator = heading
            if item.text.strip():
                output.append(Block(locator, item.text))
            for number, blip in enumerate(item._p.xpath('.//a:blip'), 1):
                from docx.oxml.ns import qn
                relation = blip.get(qn('r:embed'))
                if not relation or relation not in item.part.related_parts:
                    raise ServiceError('DOCX 包含外链或缺失图片，请嵌入图片或导出 PDF')
                data = item.part.related_parts[relation].blob
                media = image_media_type(data)
                if not media:
                    raise ServiceError('DOCX 图片格式不受支持，请转为 PNG/JPEG 或导出 PDF')
                output.append(Block(f'{locator} · 图片 {number}', image=data, media_type=media))
    walk(document)
    seen_parts = set()
    for index, section in enumerate(document.sections, 1):
        for name in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
            part = getattr(section, name)
            if part.is_linked_to_previous or part.part.partname in seen_parts:
                continue
            seen_parts.add(part.part.partname)
            walk(part, f'第 {index} 节 {name}')
    return output


def pdf_blocks(content):
    import pypdfium2 as pdfium
    from pypdfium2 import raw
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ServiceError('PDF 已加密，请解密后上传')
    if len(reader.pages) > 200:
        raise ServiceError('PDF 超过 200 页限制，请拆分文件')
    output, visual_count, image_bytes = [], 0, 0
    with PDFIUM_LOCK, closing(pdfium.PdfDocument(content)) as pdf:
        for index, source_page in enumerate(reader.pages):
            locator = f'第 {index + 1} 页'
            # Limit decompressed stream size before text extraction/layout analysis.
            stream = source_page.get_contents()
            if stream and len(stream.get_data()) > 8 * 1024 * 1024:
                raise ServiceError(f'{locator} 内容流过大，请优化 PDF 后上传')
            text = source_page.extract_text(extraction_mode='layout') or ''
            with closing(pdf[index]) as page:
                types = {obj.type for obj in page.get_objects()}
                graphical = bool(types & {raw.FPDF_PAGEOBJ_IMAGE, raw.FPDF_PAGEOBJ_PATH, raw.FPDF_PAGEOBJ_SHADING, raw.FPDF_PAGEOBJ_FORM})
                columns = sum(bool(re.search(r'\S {3,}\S', line)) for line in text.splitlines()) >= 2
                needs_vision = graphical or columns or (bool(types) and len(text.strip()) < 40) or '\ufffd' in text
                if not needs_vision:
                    if text.strip():
                        output.append(Block(locator, text))
                    continue
                visual_count += 1
                if visual_count > 20:
                    raise ServiceError('需要视觉解析的 PDF 页面超过 20 页，请按章节拆分文件')
                width, height = page.get_size()
                if min(width, height) <= 0 or max(width, height) / min(width, height) > 10:
                    raise ServiceError(f'{locator} 页面尺寸异常，请重新导出标准页面')
                scale = min(2.0, 2000 / max(width, height))
                bitmap = page.render(scale=scale)
                try:
                    image = bitmap.to_pil()
                    buffer = io.BytesIO()
                    image.save(buffer, format='PNG')
                    data = buffer.getvalue()
                finally:
                    bitmap.close()
                image_bytes += len(data)
                if image_bytes > 20 * 1024 * 1024:
                    raise ServiceError('PDF 页面图像超过 20 MB，请拆分文件')
                output.append(Block(locator + ' · 整页视觉解析', image=data, media_type='image/png'))
    return output


def extract_blocks(name, content):
    suffix = validate_upload(name, content)
    try:
        if suffix in IMAGE_SUFFIXES:
            return [Block('图片', image=content, media_type=IMAGE_MEDIA_TYPES[suffix])]
        if suffix == '.pdf':
            return pdf_blocks(content)
        if suffix == '.docx':
            return docx_blocks(content)
        text = content.decode('utf-8-sig')
        if '\x00' in text or sum(ord(c) < 32 and c not in '\n\r\t' for c in text) > 2:
            raise ServiceError('文本含二进制控制字符，请导出 UTF-8 文本后上传')
        return markdown_blocks(text) if suffix == '.md' else [Block('正文', text)]
    except ServiceError:
        raise
    except UnicodeError:
        raise ServiceError('文本资料需要使用 UTF-8 编码') from None
    except Exception:
        raise ServiceError(f'{suffix.upper().lstrip(".")} 无法解析，请修复文件或重新导出') from None


def split_document(name, content):
    blocks = extract_blocks(name, content)
    output = [chunk for block in blocks if block.text for chunk in semantic_chunks(block.locator, block.text)]
    if not output and not any(block.image for block in blocks):
        raise ServiceError('未提取到文字或图片，请补充有内容的文件')
    if len(output) > 400:
        raise ServiceError('资料超过 400 个片段限制，请拆分文件')
    return output
